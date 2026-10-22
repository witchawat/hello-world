"""Tests for loxley_sales_parser using a small synthetic workbook that mimics
the quirks of the real Loxley file (no real sales data is stored in git)."""
import datetime as dt
import sqlite3
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import loxley_sales_parser as L  # noqa: E402


def build_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # --- Sep.26: previous month, cumulative MTD + stock titled on every column
    ws = wb.create_sheet("Sep.26")
    ws.append(["", "Out of Stock"])
    ws.append(["", "Low Stock"])
    ws.append(["", "Over Stock"])
    ws.append(["Part", "Description", "LTP", "Forecast in Sep.26", "Actual", "Actual", "% ",
               "Available Stock", "Available Stock", "Cover Day", "Demand", "Promotion "])
    ws.append([None, None, None, None, "1 Sep.", " 30 Sep.", None, "31 Aug.'26", "30 Sep.'26"])
    ws.append(["006010130301", "กรีนนัท 70 ก. รสปลาหมึก", None, 100, 10, 90, 0.9, 500, 410, 12, 5, None])
    ws.append([None, "Total"])

    # --- Oct.26: current month
    ws = wb.create_sheet("Oct.26")
    ws.append([])
    ws.append([])
    ws.append([None, None, None, None, "Time Gones"])
    ws.append(["Part", "Description", "Avg Sale/mth (Y2025)", "New Weight", "Forecast in Oct.26",
               "Actual", "Actual", "Actual", "Actual", "Actual", "% ", "Plan PO in Oct.", "Pending PO.",
               "Total PO", "PO Arrival in October", None, None, "PO ขาดส่ง",
               "Available Stock", None, None, "Cover Day", "Demand", "Promotion ", "ของเข้าวันที่"])
    ws.append([None, None, None, None, None,
               "30.Sep", " 1 Oct.", "2Oct", "5 Oct.", "6 Oct.",  # 30.Sep = prior-month carry-over
               None, None, None, None, 1, 2, "Total", None,
               "1 Oct.'26", "2 Oct.'26", "5 Oct.'26"])
    ws.append(["006010130301", "กรีนนัท\xa070\xa0ก. รสปลาหมึก", 3479.37, "60 ก.", 3000,
               90, 266, 266, 497, None,  # 6 Oct not reported yet (whole column empty)
               0.16, 2300, 0, 2300, 350, None, 350, 1950,
               2858, 2703, 3011, 16.15, 15, "Tops 2 for 45.-", "วันที่ 9 ="])
    ws.append(["006010140101", "กรีนนัท 17 ก. รสมาตรฐาน", 135.3, "no change", 80,
               5, None, 6, 4, None,  # blank = 0 so far, then a return (6 -> 4)
               "#DIV/0!", 100, 0, 100, None, 100, 100, 0,
               33, 33, 126, "#DIV/0!", 18, None, None])
    ws.append([None, "Total"])

    ref = wb.create_sheet("DATA_ไซส์")
    ref.append(["Part", "Description", "ไซส์บาท", "ช่องทางที่ขาย", "ห้างที่ขาย"])
    ref.append(["006010130301", "กรีนนัท 70 ก. รสปลาหมึก", 20, None, "Makro"])
    cust = wb.create_sheet("DATA_Cocode")
    cust.append(["รหัสลูกค้า", "ชื่อลูกค้า", "Co-Code"])
    cust.append(["1377", "บริษัท ซีพี ออลล์ จำกัด (มหาชน)", "7-11"])
    cust.append(["1377", "บริษัท ซีพี ออลล์ จำกัด (มหาชน)", "7-11"])
    wb.create_sheet("Sheet1")
    wb.save(path)


@pytest.fixture()
def loaded(tmp_path):
    xlsx = tmp_path / "Green Nut Daily Stock 2026 05.10.26.xlsx"
    build_workbook(xlsx)
    db_path = tmp_path / "t.db"
    rc = L.main(["--db", f"sqlite:///{db_path}", "ingest", str(xlsx)])
    assert rc == 0
    con = sqlite3.connect(db_path)
    yield xlsx, db_path, con
    con.close()


@pytest.mark.parametrize("label,expected", [
    (" 1 Oct.", dt.date(2026, 10, 1)),
    ("2Oct", dt.date(2026, 10, 2)),
    ("1 Oct.'26", dt.date(2026, 10, 1)),
    ("30.Sep", dt.date(2026, 9, 30)),
    ("31 May.25", dt.date(2025, 5, 31)),
    ("1 June 25", dt.date(2025, 6, 1)),
    ("22Mar.", dt.date(2026, 3, 22)),
    ("Acc. 31 Mar24", dt.date(2024, 3, 31)),
    ("update 1-29 ", dt.date(2026, 10, 29)),
    (None, None),
])
def test_parse_header_date(label, expected):
    assert L.parse_header_date(label, 2026, 10) == expected


def test_header_date_year_wrap():
    assert L.parse_header_date("31.Dec", 2026, 1) == dt.date(2025, 12, 31)
    assert L.parse_header_date("02.Jan", 2025, 12) == dt.date(2026, 1, 2)


def test_sheet_month():
    assert L.sheet_month("Oct.26") == (2026, 10)
    assert L.sheet_month("June25") == (2025, 6)
    assert L.sheet_month("Jul24 ") == (2024, 7)
    assert L.sheet_month("Aug") == (None, 8)
    for bad in ("Sheet1", "forecast", "PO Jan", "DATA_Cocode"):
        with pytest.raises(ValueError):
            L.sheet_month(bad)


def test_report_date_from_name():
    assert L.report_date_from_name("Green Nut Daily Stock 2026 22.10.26.xlsx") == dt.date(2026, 10, 22)


def test_mtd_snapshots_and_daily_view(loaded):
    _, _, con = loaded
    rows = con.execute(
        "SELECT part_code, snapshot_date, mtd_qty FROM fact_sales_mtd "
        "WHERE month_start = '2026-10-01' ORDER BY part_code, snapshot_date").fetchall()
    # carry-over 30.Sep column and the unreported 6 Oct column are excluded
    assert rows == [
        ("006010130301", "2026-10-01", 266.0),
        ("006010130301", "2026-10-02", 266.0),
        ("006010130301", "2026-10-05", 497.0),
        ("006010140101", "2026-10-01", 0.0),
        ("006010140101", "2026-10-02", 6.0),
        ("006010140101", "2026-10-05", 4.0),
    ]
    daily = con.execute(
        "SELECT part_code, sales_date, sales_qty FROM v_sales_out_daily "
        "WHERE month_start = '2026-10-01' ORDER BY part_code, sales_date").fetchall()
    assert [r[2] for r in daily] == [266.0, 0.0, 231.0, 0.0, 6.0, -2.0]
    monthly = dict(con.execute(
        "SELECT substr(month_start,1,7), SUM(sales_qty) FROM v_sales_out_monthly GROUP BY 1").fetchall())
    assert monthly == {"2026-09": 90.0, "2026-10": 501.0}


def test_stock_po_and_status(loaded):
    _, _, con = loaded
    assert con.execute("SELECT stock_date, available_qty FROM fact_stock WHERE part_code='006010130301' "
                       "ORDER BY stock_date").fetchall() == [
        ("2026-08-31", 500.0), ("2026-09-30", 410.0),
        ("2026-10-01", 2858.0), ("2026-10-02", 2703.0), ("2026-10-05", 3011.0)]
    assert con.execute("SELECT part_code, arrival_date, qty FROM fact_po_arrival ORDER BY 1").fetchall() == [
        ("006010130301", "2026-10-01", 350.0), ("006010140101", "2026-10-02", 100.0)]
    st = con.execute(
        "SELECT description, forecast_qty, mtd_actual_qty, pct_of_forecast, plan_po_qty, total_po_qty, "
        "po_arrived_qty, po_outstanding_qty, available_stock, cover_days, demand, promotion, incoming_note, "
        "new_weight, avg_sale_label FROM fact_sku_status WHERE part_code='006010130301' "
        "AND report_date='2026-10-05'").fetchone()
    assert st == ("กรีนนัท 70 ก. รสปลาหมึก", 3000.0, 497.0, 0.16, 2300.0, 2300.0, 350.0, 1950.0, 3011.0,
                  16.15, 15.0, "Tops 2 for 45.-", "วันที่ 9 =", "60 ก.", "Avg Sale/mth (Y2025)")
    # Excel errors become NULL rather than breaking the load
    assert con.execute("SELECT pct_of_forecast, cover_days FROM fact_sku_status "
                       "WHERE part_code='006010140101' AND report_date='2026-10-05'").fetchone() == (None, None)


def test_dimensions(loaded):
    _, _, con = loaded
    assert con.execute("SELECT part_code, size_g, price_point, retailers FROM dim_product ORDER BY 1").fetchall() == [
        ("006010130301", 70.0, "20", "Makro"), ("006010140101", 17.0, None, None)]
    assert con.execute("SELECT * FROM dim_customer").fetchall()[0][:3] == (
        "1377", "บริษัท ซีพี ออลล์ จำกัด (มหาชน)", "7-11")


def test_reingest_is_skipped_and_force_is_idempotent(loaded):
    xlsx, db_path, con = loaded
    before = con.execute("SELECT COUNT(*) FROM fact_sales_mtd").fetchone()[0]
    db = L.DB(f"sqlite:///{db_path}")
    assert L.ingest_file(db, xlsx, False, 1, None, None, force=False)["skipped"] is True
    L.ingest_file(db, xlsx, False, 1, "msg123", "subj", force=True)
    assert con.execute("SELECT COUNT(*) FROM fact_sales_mtd").fetchone()[0] == before
    assert con.execute("SELECT email_message_id, report_date FROM ingest_log").fetchone() == ("msg123", "2026-10-05")
