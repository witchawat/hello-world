"""Incremental (diff) loading, replaying the update patterns seen in real
Loxley files from 30 Sep - 22 Oct 2026 with small synthetic workbooks."""
import calendar
import datetime as dt
import sqlite3
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import loxley_sales_parser as L  # noqa: E402

A, B = "006010130301", "006010140201"  # two SKUs
MONTH_NAMES = ["", "Jan", "Feb", "Mar", "Apr", "May", "June", "July", "Aug", "Sep", "Oct", "Nov", "Dec"]


def add_month_sheet(wb, year, month, actual, stock, po=None, total_row=True, labels=None):
    """actual / stock: {day: {sku: value}} in report order. po: {sku: {day: qty}}.
    labels: optional {day: header label} to reproduce mislabelled columns.
    Columns are laid out like Loxley's: one Actual column per report day, then
    %, Total PO, PO Arrival 1..N + Total, PO ขาดส่ง, one Available Stock column
    per day, Cover Day, Demand, Promotion. Adding a day shifts everything right."""
    po = po or {}
    days = calendar.monthrange(year, month)[1]
    ws = wb.create_sheet(f"{MONTH_NAMES[month]}.{year % 100}")
    mon = MONTH_NAMES[month][:3]
    a_days, s_days = list(actual), list(stock)
    labels = labels or {}
    row3 = [None, None, "Time Gones"] + [d / days for d in a_days]
    row4 = (["Part", "Description", "Forecast"] + ["Actual"] * len(a_days) + ["% ", "Total PO", "PO Arrival"]
            + [None] * days + ["PO ขาดส่ง"] + ["Available Stock"] * len(s_days) + ["Cover Day", "Demand", "Promotion "])
    row5 = ([None, None, None] + [labels.get(d, f" {d} {mon}.") for d in a_days] + [None, None]
            + list(range(1, days + 1)) + ["Total", None] + [f"{d} {mon}.'{year % 100}" for d in s_days])
    ws.append([])
    ws.append([])
    ws.append(row3)
    ws.append(row4)
    ws.append(row5)
    for sku in (A, B):
        arrivals = [po.get(sku, {}).get(d) for d in range(1, days + 1)]
        ws.append([sku, f"desc {sku}", 1000] + [actual[d].get(sku) for d in a_days] + [0.5, 500]
                  + arrivals + [sum(x or 0 for x in arrivals), 0]
                  + [stock[d].get(sku) for d in s_days] + [10, 3, None])
    if total_row:
        ws.append([None, "Total", None] + [sum(actual[d].get(s) or 0 for s in (A, B)) for d in a_days])
    return ws


def book(path, *months):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for m in months:
        add_month_sheet(wb, *m)
    cust = wb.create_sheet("DATA_Cocode")
    cust.append(["รหัสลูกค้า", "ชื่อลูกค้า", "Co-Code"])
    cust.append(["1377", "CP All", "7-11"])
    wb.save(path)
    return path


@pytest.fixture()
def env(tmp_path):
    db_path = tmp_path / "t.db"
    db = L.DB(f"sqlite:///{db_path}")
    db.init_schema()
    con = sqlite3.connect(db_path)
    yield tmp_path, db, con
    con.close()


def q(con, sql, *args):
    return con.execute(sql, args).fetchall()


OCT = (2026, 10)
DAY1 = {1: {A: 10, B: 5}}
DAY2 = {1: {A: 10, B: 5}, 2: {A: 30, B: 6}}
DAY5 = {1: {A: 10, B: 5}, 2: {A: 30, B: 6}, 5: {A: 70, B: 9}}
STOCK1 = {1: {A: 900, B: 90}}
STOCK2 = {1: {A: 900, B: 90}, 2: {A: 880, B: 89}}
STOCK5 = {1: {A: 900, B: 90}, 2: {A: 880, B: 89}, 5: {A: 840, B: 86}}


def test_daily_files_add_only_the_new_day(env):
    tmp, db, con = env
    r1 = L.ingest_file(db, book(tmp / "d1.xlsx", (*OCT, DAY1, STOCK1)))
    r2 = L.ingest_file(db, book(tmp / "d2.xlsx", (*OCT, DAY2, STOCK2)))
    assert (r1["sales_mtd_inserted"], r1["stock_inserted"]) == (2, 2)
    assert (r2["sales_mtd_inserted"], r2["stock_inserted"]) == (2, 2)
    assert "sales_mtd_updated" not in r2 and "sales_mtd_deleted" not in r2
    assert q(con, "SELECT sales_date, sales_qty FROM v_sales_out_daily WHERE part_code=? ORDER BY 1", A) == [
        ("2026-10-01", 10.0), ("2026-10-02", 20.0)]
    assert q(con, "SELECT COUNT(*) FROM change_log") == [(0,)]


def test_backdated_and_moved_po_arrival(env):
    tmp, db, con = env
    L.ingest_file(db, book(tmp / "d2.xlsx", (*OCT, DAY2, STOCK2, {B: {2: 350}})))
    # next day: the 350 moves from SKU B to SKU A, plus a back-dated arrival on the 1st
    r = L.ingest_file(db, book(tmp / "d5.xlsx", (*OCT, DAY5, STOCK5, {A: {1: 40, 2: 350}})))
    assert (r["po_arrival_inserted"], r["po_arrival_deleted"]) == (2, 1)
    assert q(con, "SELECT part_code, arrival_date, qty FROM fact_po_arrival ORDER BY 1, 2") == [
        (A, "2026-10-01", 40.0), (A, "2026-10-02", 350.0)]
    assert q(con, "SELECT table_name, part_code, key_date, action, old_value, file_name FROM change_log") == [
        ("fact_po_arrival", B, "2026-10-02", "delete", 350.0, "d5.xlsx")]


def test_revised_past_value_is_updated_and_logged(env):
    tmp, db, con = env
    L.ingest_file(db, book(tmp / "d2.xlsx", (*OCT, DAY2, STOCK2)))
    revised = {1: {A: 10, B: 5}, 2: {A: 28, B: 6}, 5: {A: 70, B: 9}}
    r = L.ingest_file(db, book(tmp / "d5.xlsx", (*OCT, revised, STOCK5)))
    assert r["sales_mtd_updated"] == 1
    assert q(con, "SELECT key_date, action, old_value, new_value FROM change_log") == [
        ("2026-10-02", "update", 30.0, 28.0)]


def test_mislabelled_column_fixed_from_time_gone_then_next_day(env):
    tmp, db, con = env
    # 30 Sep 2026 case: the newest column still says '1 Oct.' but holds the 2nd
    r = L.ingest_file(db, book(tmp / "bad.xlsx", (*OCT, DAY2, STOCK2, None, True, {2: " 1 Oct."})))
    assert any("using 2026-10-02 from the time-gone row" in w for w in r["warnings"])
    assert q(con, "SELECT snapshot_date, mtd_qty FROM fact_sales_mtd WHERE part_code=? ORDER BY 1", A) == [
        ("2026-10-01", 10.0), ("2026-10-02", 30.0)]
    r = L.ingest_file(db, book(tmp / "fixed.xlsx", (*OCT, DAY5, STOCK5)))
    assert r["warnings"] == [] and "sales_mtd_updated" not in r


def test_month_rollover_finalises_previous_month_then_ignores_it(env):
    tmp, db, con = env
    sep_partial = (2026, 9, {29: {A: 400, B: 40}}, {29: {A: 500, B: 50}})
    sep_final = (2026, 9, {29: {A: 400, B: 40}, 30: {A: 410, B: 41}}, {29: {A: 500, B: 50}, 30: {A: 490, B: 49}})
    L.ingest_file(db, book(tmp / "Green Nut Daily Stock 2026.xlsx", sep_partial))  # no date in the name
    r = L.ingest_file(db, book(tmp / "oct1.xlsx", sep_final, (*OCT, DAY1, STOCK1)))
    assert r["sheets"] == ["Sep.26", "Oct.26"] and r["report_date"] == "2026-10-01"
    assert q(con, "SELECT substr(month_start,1,7), sales_qty FROM v_sales_out_monthly WHERE part_code=? ORDER BY 1",
             A) == [("2026-09", 410.0), ("2026-10", 10.0)]
    # Sep.26 is re-read through the 7th (late corrections still land), then ignored
    sep_changed = (2026, 9, {29: {A: 1, B: 1}, 30: {A: 1, B: 1}}, {29: {A: 1, B: 1}, 30: {A: 1, B: 1}})
    for day in range(2, 10):
        actual = {d: {A: 10 * d, B: d} for d in range(1, day + 1)}
        stock = {d: {A: 900 - d, B: 90} for d in range(1, day + 1)}
        r = L.ingest_file(db, book(tmp / f"oct{day}.xlsx", sep_final if day <= 8 else sep_changed,
                                   (*OCT, actual, stock)))
        assert r["sheets"] == (["Sep.26", "Oct.26"] if day <= 7 else ["Oct.26"])
    assert q(con, "SELECT sales_qty FROM v_sales_out_monthly WHERE part_code=? AND month_start='2026-09-01'",
             A) == [(410.0,)]


def test_empty_next_month_sheet_is_ignored(env):
    tmp, db, _ = env
    r = L.ingest_file(db, book(tmp / "f.xlsx", (*OCT, DAY2, STOCK2), (2026, 11, {}, {})))
    assert r["sheets"] == ["Oct.26"] and r["report_date"] == "2026-10-02"


def test_duplicate_and_older_files_are_skipped(env):
    tmp, db, con = env
    L.ingest_file(db, book(tmp / "d5.xlsx", (*OCT, DAY5, STOCK5)), message_id="m5")
    assert L.ingest_file(db, book(tmp / "d5b.xlsx", (*OCT, DAY5, STOCK5)), message_id="m5")["skipped"]
    older = book(tmp / "d2.xlsx", (*OCT, DAY2, STOCK2))
    assert L.ingest_file(db, older)["skipped"].startswith("older than loaded report")
    # even when forced, an older file never deletes data after its own report date
    L.ingest_file(db, older, force=True)
    assert q(con, "SELECT COUNT(*) FROM fact_sales_mtd WHERE snapshot_date='2026-10-05'") == [(2,)]


def test_mass_delete_guard(env):
    tmp, db, con = env
    many = {d: {A: d, B: d} for d in range(1, 11)}
    L.ingest_file(db, book(tmp / "full.xlsx", (*OCT, many, {10: {A: 1, B: 1}})))
    # a broken file that lost most report columns but still reports day 10
    r = L.ingest_file(db, book(tmp / "broken.xlsx", (*OCT, {10: {A: 10, B: 10}}, {10: {A: 1, B: 1}})))
    assert any("skipped deletes" in w for w in r["warnings"])
    assert q(con, "SELECT COUNT(*) FROM fact_sales_mtd") == [(20,)]


def test_quality_warnings(env):
    tmp, db, _ = env
    path = book(tmp / "Green Nut Daily Stock 2026 06.10.26.xlsx",
                (*OCT, {1: {A: 500, B: 5}, 2: {A: 30, B: 6}}, STOCK2))
    wb = openpyxl.load_workbook(path)
    ws = wb["Oct.26"]
    ws.cell(ws.max_row, 5).value = 999  # Total row disagrees with the SKU rows
    wb.save(path)
    r = L.ingest_file(db, path, subject="Green Nut Daily Sales&Stock as of 2 Oct 2026")
    w = " | ".join(r["warnings"])
    assert "MTD fell from 500" in w
    assert "Total row says 999" in w
    assert "file name says 2026-10-06" in w
    assert "email subject" not in w  # subject matches the data


def test_read_only_load_skips_customer_master_mid_month(env):
    tmp, db, con = env
    L.ingest_file(db, book(tmp / "d1.xlsx", (*OCT, DAY1, STOCK1)))
    path = tmp / "d2.xlsx"
    book(path, (*OCT, DAY2, STOCK2))
    wb = openpyxl.load_workbook(path)
    wb["DATA_Cocode"].append(["9999", "New customer", "Makro"])
    wb.save(path)
    L.ingest_file(db, path)
    assert q(con, "SELECT customer_code FROM dim_customer") == [("1377",)]
    L.ingest_file(db, path, force=True, refresh_reference=True)
    assert q(con, "SELECT customer_code, channel_group FROM dim_customer ORDER BY 1") == [
        ("1377", "7-11"), ("9999", "Makro")]


def test_subject_date():
    assert L.report_date_from_subject("Green Nut Daily Sales&Stock as of 19 Oct 2026") == dt.date(2026, 10, 19)
    assert L.report_date_from_subject("Green Nut Daily Sales & Stock as of 30 September 2026") == dt.date(2026, 9, 30)
    assert L.report_date_from_subject(None) is None


@pytest.mark.parametrize("cmd", [["--help"], ["ingest", "--help"], ["summary", "--help"]])
def test_cli_help(cmd, capsys):
    with pytest.raises(SystemExit) as e:
        L.main(cmd)
    assert e.value.code == 0 and "usage" in capsys.readouterr().out
